#!/usr/bin/env python3
"""Gera o PGR (Programa de Gerenciamento de Risco) no modelo timbrado da Marine
a partir da proposta/apólice de RC-DC da Tokio Marine.

Princípio: nada é digitado ou resumido. Todo texto do corpo é copiado caractere
a caractere do PDF (com negrito, itálico e sublinhado), e as tabelas são lidas
célula a célula. O script só escolhe QUAIS seções entram e em que ordem.

Uso como módulo:  from pgr_core import gerar
Uso por linha de comando:
  python pgr_core.py --proposta proposta.pdf --modelo modelo.docx --saida PGR.docx
"""
import argparse
import re
import sys
import unicodedata

import pdfplumber
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

# Ordem do PGR feito à mão (Schoeler 09/2026)
SECOES_PADRAO = [
    "LIMITE MÁXIMO DE GARANTIA",
    "MERCADORIAS EXCLUÍDAS",
    "IMPORTÂNCIA SEGURADA",
    "OBJETO DO SEGURO",
    "VIAGENS",
    "PARTICIPAÇÃO OBRIGATÓRIA DO SEGURADO",
    "GERENCIAMENTO DE RISCO",
    "SINISTROS",
]

# Títulos alternativos são definidos somente por parametrização explícita.
ALIASES = {}

CORRETORA = "MARINE CORRETORA E ASSESSORIA DE SEGUROS"
TOPO_CONTEUDO = 95      # abaixo do cabeçalho da proposta Tokio
BASE_CONTEUDO = 740     # acima do rodapé "TOKIO MARINE SEGURADORA S.A."
CINZA = (0.851, 0.851, 0.851)


def norm(s):
    s = unicodedata.normalize("NFKD", s.upper())
    return re.sub(r"\s+", " ", "".join(c for c in s if not unicodedata.combining(c))).strip()


# ----------------------------------------------------------------- extração
def _dentro(obj, bbox, folga=1):
    x0, top, x1, bottom = bbox
    return obj["x0"] >= x0 - folga and obj["x1"] <= x1 + folga and \
        obj["top"] >= top - folga and obj["bottom"] <= bottom + folga


def _sublinhados(page, tabelas):
    segs = []
    for o in page.lines + page.rects:
        if abs(o["bottom"] - o["top"]) < 1.5 and (o["x1"] - o["x0"]) > 3:
            if not any(_dentro(o, t, 2) for t in tabelas):
                segs.append((o["x0"], o["x1"], o["top"]))
    return segs


def _estilo(c, segs):
    f = c["fontname"]
    sub = any(x0 - 1 <= (c["x0"] + c["x1"]) / 2 <= x1 + 1 and -0.5 <= y - c["bottom"] <= 3
              for x0, x1, y in segs)
    return ("Bold" in f, "Italic" in f or "Oblique" in f, sub)


def _runs_da_linha(chars, segs):
    """Reconstrói o texto da linha char a char, com o estilo de cada um."""
    chars = sorted(chars, key=lambda c: c["x0"])
    runs, ant = [], None
    for c in chars:
        t = c["text"]
        if ant is not None and t != " " and ant["text"] != " ":
            if c["x0"] - ant["x1"] > c["size"] * 0.18:
                t = " " + t
        st = _estilo(c, segs) if c["text"].strip() else (runs[-1][1] if runs else (False, False, False))
        if runs and runs[-1][1] == st:
            runs[-1][0] += t
        else:
            runs.append([t, st])
        ant = c
    return runs


def extrair_blocos(pdf_path):
    """Sequência de blocos: heading / linha / tabela, com posição."""
    blocos = []
    with pdfplumber.open(pdf_path) as pdf:
        for pn, page in enumerate(pdf.pages):
            cinzas = [r for r in page.rects
                      if r.get("non_stroking_color") and tuple(round(v, 3) for v in r["non_stroking_color"])[:3] == CINZA
                      and TOPO_CONTEUDO < r["top"] < BASE_CONTEUDO]
            tbls = [t for t in page.find_tables() if TOPO_CONTEUDO < t.bbox[1] < BASE_CONTEUDO]
            tb_boxes = [t.bbox for t in tbls]
            segs = _sublinhados(page, tb_boxes)
            itens = []
            for r in cinzas:
                txt = page.within_bbox((r["x0"], r["top"], r["x1"], r["bottom"])).extract_text().strip()
                if txt:
                    itens.append((r["top"], {"tipo": "titulo", "texto": re.sub(r"\s+", " ", txt), "pag": pn + 1}))
            for t in tbls:
                linhas = []
                for row in t.extract():
                    linhas.append([re.sub(r"\s*\n\s*", " ", c or "").strip() for c in row])
                # colunas nominais (para tabela com células mescladas)
                xs = sorted({round(c[0]) for c in t.cells} | {round(c[2]) for c in t.cells})
                r0 = t.rows[0].bbox
                c0 = [c for c in page.chars if _dentro(c, r0, 0.5) and c["text"].strip()]
                cab_negrito = bool(c0) and sum("Bold" in c["fontname"] for c in c0) > len(c0) / 2
                itens.append((t.bbox[1], {"tipo": "tabela", "linhas": linhas, "pag": pn + 1, "cab_negrito": cab_negrito,
                                          "bbox": t.bbox, "xs": xs}))
            # linhas de texto fora de títulos e tabelas
            excl = tb_boxes + [(r["x0"], r["top"], r["x1"], r["bottom"]) for r in cinzas]
            chars = [c for c in page.chars if TOPO_CONTEUDO < c["top"] < BASE_CONTEUDO
                     and not any(_dentro(c, b, 0.5) for b in excl)]
            chars.sort(key=lambda c: (round(c["top"]), c["x0"]))
            grupos = []
            for c in chars:
                if grupos and abs(grupos[-1][0]["top"] - c["top"]) < 2.5:
                    grupos[-1].append(c)
                else:
                    grupos.append([c])
            for g in grupos:
                texto = "".join(r[0] for r in _runs_da_linha(g, segs))
                if not texto.strip():
                    continue
                itens.append((min(c["top"] for c in g), {
                    "tipo": "linha", "runs": _runs_da_linha(g, segs), "pag": pn + 1,
                    "x0": min(c["x0"] for c in g), "x1": max(c["x1"] for c in g),
                    "top": min(c["top"] for c in g), "bottom": max(c["bottom"] for c in g)}))
            itens.sort(key=lambda i: i[0])
            blocos.extend(i[1] for i in itens)
    return blocos


INICIO_ITEM = re.compile(r"^\s*(-|•|\*|[a-z]\)|\d+(\.\d+)*\s*[-.)]|\d+\.\d+(\.\d+)*\s|[¹²³])")


def montar_paragrafos(blocos, margem_direita=565):
    """Junta linhas quebradas em parágrafos e marca espaços em branco."""
    saida, atual, ant = [], None, None
    for b in blocos:
        if b["tipo"] != "linha":
            if atual:
                saida.append(atual)
            atual, ant = None, None
            saida.append(b)
            continue
        texto = "".join(r[0] for r in b["runs"])
        gap = (b["top"] - ant["bottom"]) if (ant and ant["pag"] == b["pag"]) else 0
        continua = (atual is not None and ant["x1"] >= margem_direita and gap < 8
                    and not INICIO_ITEM.match(texto))
        if continua:
            sep = "" if atual["runs"][-1][0].endswith((" ", "-", "/")) else " "
            primeiro = b["runs"][0]
            if atual["runs"][-1][1] == primeiro[1]:
                atual["runs"][-1][0] += sep + primeiro[0]
                atual["runs"].extend([list(r) for r in b["runs"][1:]])
            else:
                atual["runs"][-1][0] += sep
                atual["runs"].extend([list(r) for r in b["runs"]])
        else:
            if atual:
                saida.append(atual)
            atual = {"tipo": "par", "runs": [list(r) for r in b["runs"]], "pag": b["pag"],
                     "espaco_antes": gap > 8}
        ant = b
    if atual:
        saida.append(atual)
    # junta tabela que continua na página seguinte
    final = []
    for b in saida:
        if (b["tipo"] == "tabela" and final and final[-1]["tipo"] == "tabela"
                and b["pag"] == final[-1]["pag"] + 1
                and len(b["linhas"][0]) == len(final[-1]["linhas"][0])):
            final[-1]["linhas"].extend(b["linhas"])
            continue
        final.append(b)
    return final


def separar_secoes(blocos):
    secoes, atual = {}, None
    ordem = []
    for b in blocos:
        if b["tipo"] == "titulo":
            atual = b["texto"]
            secoes[atual] = []
            ordem.append(atual)
        elif atual:
            secoes[atual].append(b)
    return secoes, ordem


def texto_par(b):
    return "".join(r[0] for r in b["runs"])


def dados_cabecalho(pdf_path, secoes):
    with pdfplumber.open(pdf_path) as pdf:
        linhas = pdf.pages[0].extract_text().splitlines()
    segurado = None
    for i, l in enumerate(linhas):
        if l.strip().startswith(("PROPONENTE", "INTERESSADO", "SEGURADO")) and "CNPJ" in l:
            v = linhas[i + 1]
            segurado = re.sub(r"\s+\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}.*$", "", v).strip()
            break
    vig = None
    for k, bl in secoes.items():
        if norm(k) == "VIGENCIA":
            t = " ".join(texto_par(b) for b in bl if b["tipo"] == "par")
            d = re.findall(r"\d{2}/\d{2}/\d{4}", t)
            if len(d) >= 2:
                vig = f"A PARTIR DAS 24H00 DE {d[0]} A 24H00 DE {d[1]}"
    return {"VIGÊNCIA": vig, "SEGURADO": segurado, "CORRETORA": CORRETORA, "SEGURADORA": "TOKIO MARINE"}


# ----------------------------------------------------------------- docx
def _shade(el, cor="D9D9D9"):
    pr = el.get_or_add_pPr() if hasattr(el, "get_or_add_pPr") else el.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    pr.append(shd)


def _fmt(p, depois=0, antes=0, just=True):
    pf = p.paragraph_format
    pf.space_after, pf.space_before = Pt(depois), Pt(antes)
    pf.line_spacing = 1.0
    if just:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _run(p, texto, b=False, i=False, u=False, tam=9.5):
    r = p.add_run(texto)
    r.bold, r.italic, r.underline = b or None, i or None, u or None
    r.font.name = "Arial"
    r.font.size = Pt(tam)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    return r


def _borda_tabela(tbl):
    tblPr = tbl._tbl.tblPr
    bd = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "000000")
        bd.append(e)
    tblPr.append(bd)


def gerar_docx(modelo, saida, cab, secoes, escolhidas):
    doc = Document(modelo)
    body = doc.element.body
    paras = doc.paragraphs
    # mantém: título (0), branco (1), bloco do cabeçalho (2 e 3); remove o resto do corpo
    manter = {paras[0]._p, paras[1]._p, paras[2]._p, paras[3]._p}
    for el in list(body):
        if el.tag.endswith("sectPr") or el in manter:
            continue
        body.remove(el)
    # reescreve o bloco do cabeçalho mantendo a formatação do modelo (Arial 12)
    for p, campos in ((paras[2], ("VIGÊNCIA", "SEGURADO")), (paras[3], ("CORRETORA", "SEGURADORA"))):
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        for k, campo in enumerate(campos):
            _run(p, campo, b=True, tam=12)
            _run(p, ": " + (cab[campo] or "[PREENCHER]"), tam=12)
            if k == 0:
                p.add_run().add_break()
                p.runs[-1].font.name = "Arial"
    for p in (paras[2], paras[3]):
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paras[3].paragraph_format.space_after = Pt(6)

    for nome in escolhidas:
        chave = next((k for k in secoes if norm(k) == norm(nome)), None)
        if chave is None:
            continue
        p = doc.add_paragraph()
        _fmt(p, antes=8, depois=2, just=False)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, chave, b=True)
        _shade(p._p)
        for b in secoes[chave]:
            if b["tipo"] == "par":
                if b.get("espaco_antes"):
                    _fmt(doc.add_paragraph())
                p = doc.add_paragraph()
                _fmt(p)
                for texto, (bb, ii, uu) in b["runs"]:
                    _run(p, texto, bb, ii, uu)
            elif b["tipo"] == "tabela":
                linhas = b["linhas"]
                ncol = max(len(l) for l in linhas)
                t = doc.add_table(rows=len(linhas), cols=ncol)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                _borda_tabela(t)
                for ri, row in enumerate(linhas):
                    for ci in range(ncol):
                        val = row[ci] if ci < len(row) else ""
                        cell = t.cell(ri, ci)
                        cp = cell.paragraphs[0]
                        _fmt(cp, just=False)
                        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        _run(cp, val or "", b=(ri == 0 and b.get("cab_negrito", True)), tam=8.5)
                sp = doc.add_paragraph(); _fmt(sp)
    doc.save(saida)


# ----------------------------------------------------------------- conferência
def conferir(pdf_path, secoes, escolhidas, docx_path):
    """Confere que cada palavra do PDF (nas seções escolhidas) está no .docx, em ordem."""
    def palavras_pdf():
        out = []
        for nome in escolhidas:
            chave = next((k for k in secoes if norm(k) == norm(nome)), None)
            if not chave:
                continue
            out += chave.split()
            for b in secoes[chave]:
                if b["tipo"] == "par":
                    out += texto_par(b).split()
                else:
                    for row in b["linhas"]:
                        for c in row:
                            out += (c or "").split()
        return out
    d = Document(docx_path)
    w_doc = []
    for el in d.element.body.iter():
        if el.tag == qn("w:p"):
            w_doc += "".join(t.text or "" for t in el.iter(qn("w:t"))).split()
    a, b = palavras_pdf(), w_doc
    # comparação de sequência
    import difflib
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    problemas = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag in ("replace", "delete"):
            problemas.append(f"{tag}: proposta='{' '.join(a[i1:i2])[:120]}' docx='{' '.join(b[j1:j2])[:120]}'")
    return len(a), problemas


def listar_secoes(proposta):
    blocos = montar_paragrafos(extrair_blocos(proposta))
    return separar_secoes(blocos)[1]


def resolver_secao(nome, secoes, mapa=None):
    """Acha no PDF o título que corresponde à seção padrão `nome`.
    mapa[nome] pode ser um título, uma lista de títulos ou "__PULAR__"."""
    mapa = mapa or {}
    preferidos = mapa.get(nome)
    if preferidos == "__PULAR__":
        return "__PULAR__"
    if isinstance(preferidos, str):
        preferidos = [preferidos]
    for cand in list(preferidos or []) + ALIASES.get(nome, []) + [nome]:
        for k in secoes:
            if norm(k) == norm(cand):
                return k
    return None


def gerar(proposta, modelo, saida, secoes_escolhidas=None, mapa=None):
    """Gera o PGR e devolve um relatório (dict). relatorio['ok'] só é True
    quando o cabeçalho foi todo encontrado, todas as seções existem e a
    conferência literal deu zero divergência."""
    escolhidas = list(secoes_escolhidas or SECOES_PADRAO)
    with pdfplumber.open(proposta) as pdf:
        tem_texto = sum(len(p.chars) for p in pdf.pages[:3]) > 50
    if not tem_texto:
        return {"ok": False, "erro": ("Este PDF não tem texto, só desenho das letras. Isso acontece quando o arquivo "
                                      "passa por 'Microsoft Print to PDF' ou é escaneado (comum na versão assinada). "
                                      "Use o PDF original baixado do portal da Tokio, sem imprimir/assinar."),
                "cabecalho": {}, "incluidas": [], "faltando": escolhidas, "puladas": [], "fora": [], "palavras": 0, "problemas": []}
    blocos = montar_paragrafos(extrair_blocos(proposta))
    secoes, ordem = separar_secoes(blocos)
    if not secoes:
        return {"ok": False, "erro": "Nenhuma seção reconhecida. O layout não parece ser de proposta/apólice Tokio Marine.",
                "cabecalho": {}, "incluidas": [], "faltando": escolhidas, "puladas": [], "fora": [], "palavras": 0, "problemas": []}
    cab = dados_cabecalho(proposta, secoes)
    mapa = mapa or {}
    resolvidas, faltando, puladas = [], [], []
    for nome in escolhidas:
        chave = resolver_secao(nome, secoes, mapa)
        if chave == "__PULAR__":
            puladas.append(nome)
        elif chave:
            resolvidas.append((nome, chave))
        else:
            faltando.append(nome)
    chaves = [c for _, c in resolvidas]
    gerar_docx(modelo, saida, cab, secoes, chaves)
    n, probs = conferir(proposta, secoes, chaves, saida)
    sem_cab = [k for k, v in cab.items() if not v]
    usadas = {norm(c) for c in chaves}
    return {
        "ok": not probs and not faltando and not sem_cab,
        "erro": None,
        "cabecalho": cab,
        "cabecalho_faltando": sem_cab,
        "incluidas": [nm for nm, c in resolvidas],
        "faltando": faltando,
        "puladas": puladas,
        "fora": [k for k in ordem if norm(k) not in usadas],
        "secoes_pdf": ordem,
        "palavras": n,
        "problemas": probs,
        "saida": saida,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--proposta", required=True)
    ap.add_argument("--modelo", required=True)
    ap.add_argument("--saida", required=True)
    ap.add_argument("--secoes", default=",".join(SECOES_PADRAO))
    ap.add_argument("--listar", action="store_true", help="só lista as seções encontradas")
    a = ap.parse_args()
    if a.listar:
        print("\n".join(listar_secoes(a.proposta))); return
    r = gerar(a.proposta, a.modelo, a.saida, [s.strip() for s in a.secoes.split(",") if s.strip()])
    if r["erro"]:
        print(r["erro"]); sys.exit(2)
    print("CABEÇALHO:")
    for k, v in r["cabecalho"].items():
        print(f"  {k}: {v or '[NÃO ENCONTRADO - PREENCHER]'}")
    print(f"SEÇÕES INCLUÍDAS: {r['incluidas']}")
    if r["faltando"]:
        print(f"SEÇÕES NÃO ENCONTRADAS NA PROPOSTA: {r['faltando']}")
    print(f"SEÇÕES DA PROPOSTA QUE FICARAM DE FORA: {r['fora']}")
    print(f"CONFERÊNCIA LITERAL: {r['palavras']} palavras comparadas, {len(r['problemas'])} divergência(s)")
    for p in r["problemas"][:30]:
        print("  - " + p)
    sys.exit(0 if r["ok"] else 1)


if __name__ == "__main__":
    main()
